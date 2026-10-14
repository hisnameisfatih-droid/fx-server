"""Staatsanleihe-Renditen aus offiziellen Quellen (Tagesschlusswerte, teils 1 Tag Verzug).

- US 2y/10y:   US-Finanzministerium (Treasury Daily Par Yield Curve, CSV)
- JGB 2y/10y:  Japanisches Finanzministerium (MoF, CSV)
- Gilt 2y/10y/30y: Bank of England (Nominal-Zinskurve, Excel im ZIP); Fallback Gilt 10y: BoE-Datenbank
"""
from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime, timedelta

from ..common import http_get

SERIEN = {
    "uk10": "Gilt 10y", "uk30": "Gilt 30y", "jp10": "JGB 10y", "jp2": "JGB 2y", "us10": "US 10y",
    "uk2": "Gilt 2y", "us2": "US 2y",
}
KERN = ("uk10", "uk30", "jp10", "jp2", "us10")   # fehlen diese, steht es im Report; 2y-Reihen sind Zusatz


def _f(x) -> float | None:
    try:
        v = float(str(x).strip())
        return v if v == v else None  # NaN raus
    except (TypeError, ValueError):
        return None


# ---------------- USA ----------------
def us_treasury(jahre: list[int] | None = None) -> dict[str, dict[str, float]]:
    jahre = jahre or [date.today().year]
    erg: dict[str, dict[str, float]] = {}
    for jahr in jahre:
        url = (f"https://home.treasury.gov/resource-center/data-chart-center/interest-rates/"
               f"daily-treasury-rates.csv/{jahr}/all?type=daily_treasury_yield_curve"
               f"&field_tdr_date_value={jahr}&page&_format=csv")
        erg.update(parse_treasury(http_get(url, timeout=30).text))
    return erg


def parse_treasury(text: str) -> dict[str, dict[str, float]]:
    erg = {}
    for row in csv.DictReader(io.StringIO(text)):
        try:
            d = datetime.strptime(row["Date"], "%m/%d/%Y").date().isoformat()
        except (KeyError, ValueError):
            continue
        for schluessel, spalte in (("us10", "10 Yr"), ("us2", "2 Yr")):
            v = _f(row.get(spalte))
            if v is not None:
                erg.setdefault(d, {})[schluessel] = v
    return erg


# ---------------- Japan ----------------
def jgb(mit_historie: bool = False) -> dict[str, dict[str, float]]:
    erg = parse_mof(http_get("https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/jgbcme.csv",
                             timeout=30).text)
    if mit_historie:
        try:
            alt = parse_mof(http_get(
                "https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/historical/jgbcme_all.csv",
                timeout=60).text, ab=date.today() - timedelta(days=90))
            alt.update(erg)
            erg = alt
        except Exception:  # noqa: BLE001
            pass
    return erg


def parse_mof(text: str, ab: date | None = None) -> dict[str, dict[str, float]]:
    erg = {}
    zeilen = text.splitlines()
    kopf_idx = next((i for i, z in enumerate(zeilen) if z.startswith("Date,")), None)
    if kopf_idx is None:
        raise ValueError("MoF-CSV ohne Kopfzeile")
    for row in csv.DictReader(io.StringIO("\n".join(zeilen[kopf_idx:]))):
        try:
            d = datetime.strptime(row["Date"].strip(), "%Y/%m/%d").date()
        except (KeyError, ValueError):
            continue
        if ab and d < ab:
            continue
        werte = {k: _f(row.get(spalte)) for k, spalte in (("jp10", "10Y"), ("jp2", "2Y"))}
        werte = {k: v for k, v in werte.items() if v is not None}
        if werte:
            erg[d.isoformat()] = werte
    return erg


# ---------------- Großbritannien ----------------
BOE_ZIP = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/latest-yield-curve-data.zip"


def gilts() -> dict[str, dict[str, float]]:
    """Gilt 2y/10y/30y (Nominal-Spotkurve). Fallback 10y: BoE-Datenbank (Par Yield)."""
    try:
        erg = parse_boe_zip(http_get(BOE_ZIP, timeout=60).content)
        if erg:
            return erg
    except Exception:  # noqa: BLE001
        pass
    return boe_iadb_10y()


BOE_LAUFZEITEN = (("uk2", 2.0), ("uk10", 10.0), ("uk30", 30.0))


def parse_boe_zip(inhalt: bytes) -> dict[str, dict[str, float]]:
    """Liest alle Spotkurven-Blätter. Die BoE-Datei hat ein Blatt 'spot, short end' (Laufzeiten in Monaten)
    vor dem Blatt 'spot curve' (Jahre) – beide werden gelesen, Monate werden in Jahre umgerechnet."""
    from openpyxl import load_workbook

    erg: dict[str, dict[str, float]] = {}
    with zipfile.ZipFile(io.BytesIO(inhalt)) as z:
        namen = [n for n in z.namelist() if "nominal" in n.lower() and "daily" in n.lower() and n.endswith(".xlsx")]
        for name in namen:
            wb = load_workbook(io.BytesIO(z.read(name)), read_only=True, data_only=True)
            for blatt_name in [s for s in wb.sheetnames if "spot" in s.lower()]:
                for d, werte in _spot_blatt(wb[blatt_name]).items():
                    ziel = erg.setdefault(d, {})
                    for k, v in werte.items():
                        ziel.setdefault(k, v)   # erstes Blatt mit dem Wert gewinnt
    return erg


def _spot_blatt(blatt) -> dict[str, dict[str, float]]:
    erg: dict[str, dict[str, float]] = {}
    laufzeiten: dict[int, float] = {}   # Spalte -> Laufzeit in Jahren
    for row in blatt.iter_rows(values_only=True):
        if not row:
            continue
        kopf = row[0]
        if not laufzeiten and isinstance(kopf, str) and ("year" in kopf.lower() or "month" in kopf.lower()):
            faktor = 1.0 if "year" in kopf.lower() else 1 / 12
            for i, v in enumerate(row[1:], start=1):
                f = _f(v)
                if f is not None:
                    laufzeiten[i] = f * faktor
            continue
        if laufzeiten and isinstance(kopf, (datetime, date)):
            d = (kopf.date() if isinstance(kopf, datetime) else kopf).isoformat()
            for schluessel, jahre in BOE_LAUFZEITEN:
                spalte = next((i for i, j in laufzeiten.items() if abs(j - jahre) < 1e-6), None)
                if spalte is not None and spalte < len(row):
                    v = _f(row[spalte])
                    if v is not None:
                        erg.setdefault(d, {})[schluessel] = v
    return erg


def boe_iadb_10y() -> dict[str, dict[str, float]]:
    von = (date.today() - timedelta(days=60)).strftime("%d/%b/%Y")
    url = ("https://www.bankofengland.co.uk/boeapps/database/_iadb-fromshowcolumns.asp"
           f"?csv.x=yes&Datefrom={von}&Dateto=now&SeriesCodes=IUDMNPY&CSVF=TN&UsingCodes=Y&VPD=Y&VFD=N")
    return parse_iadb(http_get(url, timeout=30).text)


def parse_iadb(text: str) -> dict[str, dict[str, float]]:
    erg = {}
    for row in csv.reader(io.StringIO(text)):
        if len(row) < 2:
            continue
        try:
            d = datetime.strptime(row[0].strip(), "%d %b %Y").date().isoformat()
        except ValueError:
            continue
        v = _f(row[1])
        if v is not None:
            erg[d] = {"uk10": v}
    return erg


def lade_alle(mit_historie: bool = False) -> tuple[dict[str, dict[str, float]], list[str]]:
    """Gibt ({datum: {serie: wert}}, fehlende Quellen) zurück."""
    alle: dict[str, dict[str, float]] = {}
    fehler = []
    jahre = [date.today().year] + ([date.today().year - 1] if mit_historie or date.today().month == 1 else [])
    for name, fn in (("US-Treasury", lambda: us_treasury(jahre)),
                     ("JGB (MoF)", lambda: jgb(mit_historie)),
                     ("Gilts (BoE)", gilts)):
        try:
            teil = fn()
            if not teil:
                raise ValueError("leer")
            for d, werte in teil.items():
                alle.setdefault(d, {}).update(werte)
        except Exception:  # noqa: BLE001
            fehler.append(name)
    return alle, fehler
